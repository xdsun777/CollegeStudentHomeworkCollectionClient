#!/usr/bin/env python3
"""
网盘作业管理工具
功能：
1. 全量同步网盘所有文件到本地 WORKDIR
2. 同步指定学科或学科下指定作业（每个学生）
3. 打包指定作业
4. 生成报告
5. 创建/删除作业文件夹
"""

import os
import sys
import shutil
import argparse
import requests
import zipfile
import re
import base64
from io import BytesIO
from pathlib import Path
from dotenv import load_dotenv
from urllib.parse import quote

load_dotenv()

WORKDIR = Path(os.getenv('WORKDIR', ''))
BASE_URL = os.getenv('BASE_URL', '').rstrip('/')
USER = os.getenv('USER', '')
PASSWD = os.getenv('PASSWD', '')

session = requests.Session()
# session.auth = (USER, PASSWD)
auth_str = f"{USER}:{PASSWD}"
auth_bytes = auth_str.encode('utf-8')
auth_b64 = base64.b64encode(auth_bytes).decode('ascii')
session.headers.update({'Authorization': f'Basic {auth_b64}'})

def encode_path(path: str) -> str:
    """对路径进行URL编码，保留斜杠"""
    if not path:
        return ''
    return '/'.join(quote(part, safe='') for part in path.split('/'))

def ensure_dir(path: Path):
    path.mkdir(parents=True, exist_ok=True)

def is_student_dir(name: str) -> bool:
    return bool(re.match(r'^\d+-\w+$', name))

def list_remote_dir(remote_path: str):
    url = f"{BASE_URL}/{encode_path(remote_path)}".rstrip('/') + '/?json'
    try:
        resp = session.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        entries = []
        for item in data.get('paths', []):
            name = item['name']
            entries.append({
                'name': name,
                'type': 'directory' if item['path_type'] == 'Dir' else 'file',
                'size': item.get('size', 0)
            })
        return entries
    except Exception as e:
        print(f"无法列出远程目录 {remote_path}: {e}")
        return []

def create_remote_dir(remote_path: str):
    url = f"{BASE_URL}/{encode_path(remote_path)}"
    try:
        resp = session.request('MKCOL', url, timeout=10)
        if resp.status_code in (200, 201):
            return True
        elif resp.status_code == 409:
            return False
        else:
            print(f"  创建目录失败 {remote_path}: {resp.status_code}")
            return False
    except Exception as e:
        print(f"  创建目录异常 {remote_path}: {e}")
        return False

def delete_remote(remote_path: str):
    url = f"{BASE_URL}/{encode_path(remote_path)}"
    try:
        resp = session.delete(url, timeout=10)
        if resp.status_code in (200, 204):
            return True
        elif resp.status_code == 404:
            return False
        else:
            print(f"  删除失败 {remote_path}: {resp.status_code}")
            return False
    except Exception as e:
        print(f"  删除异常 {remote_path}: {e}")
        return False

def download_file(remote_path: str, local_path: Path, force=False):
    if not force and local_path.exists():
        try:
            head = session.head(f"{BASE_URL}/{remote_path}", timeout=5)
            remote_size = int(head.headers.get('Content-Length', 0))
            if remote_size == local_path.stat().st_size:
                print(f"  已存在，跳过：{local_path}")
                return
        except:
            pass

    url = f"{BASE_URL}/{encode_path(remote_path)}"
    try:
        print(f"  下载: {url} -> {local_path}")
        resp = session.get(url, stream=True, timeout=30)
        resp.raise_for_status()
        ensure_dir(local_path.parent)
        with open(local_path, 'wb') as f:
            for chunk in resp.iter_content(chunk_size=8192):
                if chunk:
                    f.write(chunk)
    except Exception as e:
        print(f"  下载失败 {remote_path}: {e}")

# def download_zip(remote_dir: str, local_dir: Path, force=False):
#     if not force and local_dir.exists() and any(local_dir.iterdir()):
#         print(f"  本地目录已存在且非空，跳过：{local_dir}")
#         return

#     url = f"{BASE_URL}/{remote_dir}?zip"
#     try:
#         print(f"  下载并解压: {url} -> {local_dir}")
#         resp = session.get(url, timeout=30)
#         resp.raise_for_status()
#         ensure_dir(local_dir)
#         with zipfile.ZipFile(BytesIO(resp.content)) as zf:
#             import tempfile
#             with tempfile.TemporaryDirectory() as tmpdir:
#                 tmp_path = Path(tmpdir)
#                 zf.extractall(tmp_path)
#                 extracted = list(tmp_path.iterdir())
#                 if len(extracted) == 1 and extracted[0].is_dir():
#                     subdir = extracted[0]
#                     for item in subdir.iterdir():
#                         dest = local_dir / item.name
#                         if dest.exists() and force:
#                             if dest.is_dir():
#                                 import shutil
#                                 shutil.rmtree(dest)
#                             else:
#                                 dest.unlink()
#                         item.rename(dest)
#                 else:
#                     for item in extracted:
#                         dest = local_dir / item.name
#                         if dest.exists() and force:
#                             if dest.is_dir():
#                                 import shutil
#                                 shutil.rmtree(dest)
#                             else:
#                                 dest.unlink()
#                         item.rename(dest)
#     except Exception as e:
#         print(f"  下载失败 {remote_dir}: {e}")


def download_zip(remote_dir: str, local_dir: Path, force=False):
    if not force and local_dir.exists() and any(local_dir.iterdir()):
        print(f"  本地目录已存在且非空，跳过：{local_dir}")
        return

    url = f"{BASE_URL}/{remote_dir}?zip"
    try:
        print(f"  下载并解压: {url} -> {local_dir}")
        resp = session.get(url, timeout=30)
        resp.raise_for_status()
        ensure_dir(local_dir)
        with zipfile.ZipFile(BytesIO(resp.content)) as zf:
            import tempfile
            with tempfile.TemporaryDirectory() as tmpdir:
                tmp_path = Path(tmpdir)
                zf.extractall(tmp_path)
                extracted = list(tmp_path.iterdir())
                if len(extracted) == 1 and extracted[0].is_dir():
                    subdir = extracted[0]
                    for item in subdir.iterdir():
                        dest = local_dir / item.name
                        if dest.exists() and force:
                            if dest.is_dir():
                                shutil.rmtree(dest)
                            else:
                                dest.unlink()
                        shutil.move(item, dest)
                else:
                    for item in extracted:
                        dest = local_dir / item.name
                        if dest.exists() and force:
                            if dest.is_dir():
                                shutil.rmtree(dest)
                            else:
                                dest.unlink()
                        shutil.move(item, dest)
    except Exception as e:
        print(f"  下载失败 {remote_dir}: {e}")


def sync_all(force=False):
    if not WORKDIR:
        print("错误：未设置 WORKDIR 环境变量，无法执行全量同步。")
        return

    print(f"开始全量同步网盘到 {WORKDIR}")

    def _sync_recursive(remote_path: str, local_path: Path):
        entries = list_remote_dir(remote_path)
        for entry in entries:
            name = entry['name']
            if entry['type'] == 'directory':
                new_remote = f"{remote_path}/{name}" if remote_path else name
                new_local = local_path / name
                ensure_dir(new_local)
                _sync_recursive(new_remote, new_local)
            else:
                file_remote = f"{remote_path}/{name}" if remote_path else name
                file_local = local_path / name
                download_file(file_remote, file_local, force)

    _sync_recursive('', WORKDIR)
    print("全量同步完成。")

# ========== 新增：同步学科/作业 ==========
def sync_assignment(course: str, assignment: str = None, force=False):
    """
    为每个学生同步指定学科下的所有作业（或指定作业）
    course: 学科名称
    assignment: 作业名称，如果为 None 则同步该学科下所有作业
    force: 强制覆盖
    """
    if not WORKDIR:
        print("错误：未设置 WORKDIR 环境变量。")
        return

    # 获取所有学生
    try:
        local_students = [d.name for d in WORKDIR.iterdir() if d.is_dir() and is_student_dir(d.name)]
    except Exception as e:
        print(f"无法读取本地学生目录：{e}")
        local_students = []

    if not local_students:
        print("本地 WORKDIR 中未找到学生文件夹，将尝试从网盘获取学生列表。")
        remote_root = list_remote_dir('')
        remote_students = [e['name'] for e in remote_root if e['type'] == 'directory' and is_student_dir(e['name'])]
        if not remote_students:
            print("网盘根目录下也没有找到学生文件夹。")
            return
        students = remote_students
    else:
        students = local_students

    print(f"共 {len(students)} 个学生。")

    for student in students:
        if assignment is None:
            # 同步该学科下的所有作业：下载整个学科目录
            remote_course_dir = f"{student}/{course}"
            local_course_dir = WORKDIR / student / course
            if not force and local_course_dir.exists() and any(local_course_dir.iterdir()):
                print(f"  跳过 {student} 的 {course} 目录（已存在且非空）")
                continue
            print(f"同步 {student} 的 {course} 目录...")
            download_zip(remote_course_dir, local_course_dir, force)
        else:
            # 只同步指定作业
            remote_job_dir = f"{student}/{course}/{assignment}"
            local_job_dir = WORKDIR / student / course / assignment
            if not force and local_job_dir.exists() and any(local_job_dir.iterdir()):
                print(f"  跳过 {student} 的 {course}/{assignment}（已存在且非空）")
                continue
            print(f"同步 {student} 的 {course}/{assignment}...")
            download_zip(remote_job_dir, local_job_dir, force)

    print("同步完成。")

# ========== 以下为原有函数（打包、报告、创建、删除）保持不变 ==========
def fetch_assignment(course: str, assignment: str, force=False):
    """确保本地 WORKDIR 中存在所有学生的该作业（若缺失则从网盘下载）"""
    if not WORKDIR:
        print("错误：未设置 WORKDIR 环境变量。")
        return None

    try:
        local_students = [d.name for d in WORKDIR.iterdir() if d.is_dir() and is_student_dir(d.name)]
    except Exception as e:
        print(f"无法读取本地学生目录：{e}")
        local_students = []

    if not local_students:
        print("本地 WORKDIR 中未找到学生文件夹，将尝试从网盘获取学生列表。")
        remote_root = list_remote_dir('')
        remote_students = [e['name'] for e in remote_root if e['type'] == 'directory' and is_student_dir(e['name'])]
        if not remote_students:
            print("网盘根目录下也没有找到学生文件夹。")
            return None
        students = remote_students
    else:
        students = local_students

    print(f"共 {len(students)} 个学生。")
    for student in students:
        local_job_dir = WORKDIR / student / course / assignment
        remote_job_dir = f"{student}/{course}/{assignment}"
        need_download = force
        if not need_download:
            if not local_job_dir.exists() or not any(local_job_dir.iterdir()):
                need_download = True
        if need_download:
            print(f"下载 {student} 的作业...")
            download_zip(remote_job_dir, local_job_dir, force)

    return WORKDIR

def zip_assignment(course: str, assignment: str, output_zip: Path, force=False):
    """将本地所有学生的指定作业打包成 ZIP"""
    if not WORKDIR:
        print("错误：未设置 WORKDIR 环境变量。")
        return

    fetch_assignment(course, assignment, force)

    students = []
    try:
        for d in WORKDIR.iterdir():
            if d.is_dir() and is_student_dir(d.name):
                students.append(d)
    except Exception as e:
        print(f"无法读取本地学生目录：{e}")

    if not students:
        print("没有找到任何学生文件夹，无法打包。")
        return

    ensure_dir(output_zip.parent)
    with zipfile.ZipFile(output_zip, 'w', zipfile.ZIP_DEFLATED) as zf:
        for student_dir in students:
            student_name = student_dir.name
            job_dir = student_dir / course / assignment
            if not job_dir.exists() or not any(job_dir.iterdir()):
                print(f"跳过 {student_name}：无此作业")
                continue

            for file_path in job_dir.rglob('*'):
                if file_path.is_file():
                    rel_path = file_path.relative_to(job_dir)
                    arcname = f"{student_name}/{rel_path}"
                    zf.write(file_path, arcname)
            print(f"已添加 {student_name} 的作业")
    print(f"打包完成：{output_zip}")

def generate_report(output_excel: Path):
    """生成按科目分工作表的完成情况报告，√ 表示完成"""
    if not WORKDIR:
        print("错误：未设置 WORKDIR 环境变量。")
        return

    student_data = {}
    courses = {}

    student_dirs = [d for d in WORKDIR.iterdir() if d.is_dir() and is_student_dir(d.name)]
    if not student_dirs:
        print("没有找到任何学生文件夹，无法生成报告。")
        return

    for student_dir in student_dirs:
        student_name = student_dir.name
        student_data[student_name] = {}
        for course_dir in student_dir.iterdir():
            if not course_dir.is_dir():
                continue
            course_name = course_dir.name
            student_data[student_name][course_name] = {}
            if course_name not in courses:
                courses[course_name] = set()
            for assignment_dir in course_dir.iterdir():
                if not assignment_dir.is_dir():
                    continue
                assignment_name = assignment_dir.name
                has_files = any(assignment_dir.iterdir())
                student_data[student_name][course_name][assignment_name] = has_files
                courses[course_name].add(assignment_name)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment
    except ImportError:
        print("请安装 openpyxl: pip install openpyxl")
        return

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    for course_name, assignments in sorted(courses.items()):
        if not assignments:
            continue
        ws = wb.create_sheet(title=course_name)
        headers = ["学号", "姓名"] + sorted(assignments)
        ws.append(headers)

        header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for student_name in sorted(student_data.keys()):
            parts = student_name.split('-', 1)
            student_id = parts[0] if len(parts) > 0 else ''
            student_real_name = parts[1] if len(parts) > 1 else ''
            row = [student_id, student_real_name]
            for assignment in sorted(assignments):
                completed = student_data.get(student_name, {}).get(course_name, {}).get(assignment, False)
                row.append("√" if completed else "")
            ws.append(row)

        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[col_letter].width = adjusted_width

    if not courses:
        ws = wb.create_sheet(title="无数据")
        ws.append(["提示"])
        ws.append(["未找到任何科目或作业文件夹"])

    ensure_dir(output_excel.parent)
    wb.save(output_excel)
    print(f"报告已生成：{output_excel}")

def new_assignment(course: str, assignment: str):
    """为所有学生在远程创建指定作业文件夹"""
    print("正在获取学生列表...")
    remote_root = list_remote_dir('')
    students = [e['name'] for e in remote_root if e['type'] == 'directory' and is_student_dir(e['name'])]
    if not students:
        print("未找到任何学生文件夹。")
        return

    print(f"共 {len(students)} 个学生。")
    for student in students:
        course_path = f"{student}/{course}"
        print(f"处理 {student}...")
        if not create_remote_dir(course_path):
            pass
        assignment_path = f"{course_path}/{assignment}"
        if create_remote_dir(assignment_path):
            print(f"  已创建: {assignment_path}")
        else:
            print(f"  目录已存在: {assignment_path}")
    print("完成。")

def delete_assignment(course: str, assignment: str, yes=False):
    """删除所有学生的指定作业文件夹（远程）"""
    print("正在获取学生列表...")
    remote_root = list_remote_dir('')
    students = [e['name'] for e in remote_root if e['type'] == 'directory' and is_student_dir(e['name'])]
    if not students:
        print("未找到任何学生文件夹。")
        return

    paths_to_delete = []
    for student in students:
        job_path = f"{student}/{course}/{assignment}"
        url = f"{BASE_URL}/{job_path}"
        try:
            resp = session.head(url, timeout=5)
            if resp.status_code == 200:
                paths_to_delete.append(job_path)
        except:
            pass

    if not paths_to_delete:
        print("没有找到任何要删除的作业文件夹。")
        return

    print(f"将删除以下 {len(paths_to_delete)} 个作业文件夹：")
    for p in paths_to_delete:
        print(f"  {p}")

    if not yes:
        confirm = input("确认删除以上所有文件夹及其内容？(y/N): ").strip().lower()
        if confirm not in ('y', 'yes'):
            print("取消删除。")
            return

    print("开始删除...")
    for p in paths_to_delete:
        if delete_remote(p):
            print(f"  已删除: {p}")
        else:
            print(f"  删除失败: {p}")
    print("删除完成。")

def main():
    parser = argparse.ArgumentParser(description="网盘作业管理工具")
    subparsers = parser.add_subparsers(dest='command', help='子命令')

    # 全量同步模式（可选指定学科/作业）
    sync_parser = subparsers.add_parser('sync', help='同步网盘到本地')
    sync_parser.add_argument('--force', '-f', action='store_true', help='强制覆盖已存在文件')
    sync_parser.add_argument('-c', '--course', help='学科名称（可选，不填则全量同步）')
    sync_parser.add_argument('-a', '--assignment', help='作业名称（可选，需与 -c 配合使用）')

    # 打包模式
    pack_parser = subparsers.add_parser('pack', help='打包指定作业')
    pack_parser.add_argument('-c', '--course', required=True, help='科目名称')
    pack_parser.add_argument('-a', '--assignment', required=True, help='作业名称')
    pack_parser.add_argument('-t', '--target-zip', required=True, help='输出 ZIP 文件路径')
    pack_parser.add_argument('--force', '-f', action='store_true', help='强制覆盖已存在文件')

    # 报告模式
    report_parser = subparsers.add_parser('report', help='生成统计报告')
    report_parser.add_argument('-r', '--report', required=True, help='输出 Excel 文件路径')

    # 新建作业文件夹模式
    new_parser = subparsers.add_parser('new', help='为所有学生创建作业文件夹')
    new_parser.add_argument('-c', '--course', required=True, help='科目名称')
    new_parser.add_argument('-a', '--assignment', required=True, help='作业名称')

    # 删除作业文件夹模式
    delete_parser = subparsers.add_parser('delete', help='删除所有学生的指定作业文件夹')
    delete_parser.add_argument('-c', '--course', required=True, help='科目名称')
    delete_parser.add_argument('-a', '--assignment', required=True, help='作业名称')
    delete_parser.add_argument('--yes', action='store_true', help='跳过确认直接删除')

    args = parser.parse_args()

    if args.command == 'sync':
        if args.course:
            # 如果指定了学科，则同步该学科（若同时指定作业则只同步该作业）
            sync_assignment(args.course, args.assignment, force=args.force)
        else:
            sync_all(force=args.force)
    elif args.command == 'pack':
        zip_assignment(args.course, args.assignment, Path(args.target_zip), args.force)
    elif args.command == 'report':
        generate_report(Path(args.report))
    elif args.command == 'new':
        new_assignment(args.course, args.assignment)
    elif args.command == 'delete':
        delete_assignment(args.course, args.assignment, yes=args.yes)
    else:
        # 兼容旧版无子命令的调用
        if hasattr(args, 'course') and args.course:
            # 如果传入了 -c，则视为同步学科/作业
            sync_assignment(args.course, getattr(args, 'assignment', None), force=getattr(args, 'force', False))
        elif hasattr(args, 'report') and args.report:
            generate_report(Path(args.report))
        else:
            sync_all(force=getattr(args, 'force', False))

if __name__ == '__main__':
    main()